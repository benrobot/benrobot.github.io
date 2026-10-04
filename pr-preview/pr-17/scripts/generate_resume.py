#!/usr/bin/env python3
"""Build benjamin-roedell-resume.docx from the site's resume data.

Email and phone are optional and must come from outside this repository.
Without them, the public resume is unchanged.
"""

import argparse
import os
import re
from html import unescape
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FONT = "Calibri"
INK = RGBColor(0x1C, 0x1F, 0x24)


def load_yaml(path):
    with path.open(encoding="utf-8") as handle:
        return yaml.safe_load(handle)


def strip_tags(value):
    text = unescape(str(value or ""))
    text = re.sub(r"(?is)<br\s*/?>", " ", text)
    text = re.sub(r"(?is)<[^>]+>", "", text)
    return re.sub(r"\s+", " ", text).strip()


def html_blocks(value):
    text = unescape(str(value or "")).replace("\r", "")
    text = re.sub(r"(?i)<br\s*/?>\s*<br\s*/?>", "\n\n", text)
    text = re.sub(r"(?i)<br\s*/?>", "\n", text)
    blocks = []
    for chunk in re.split(r"\n\s*\n", text):
        chunk = chunk.strip()
        if not chunk:
            continue
        match = re.match(r"(?is)<strong>(.*?)</strong>\s*:?\s*(.*)", chunk)
        if match:
            label = strip_tags(match.group(1))
            body = strip_tags(match.group(2))
            blocks.append((label, body))
        else:
            blocks.append((None, strip_tags(chunk)))
    return blocks


def set_run_font(run, size, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = INK
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_fonts.set(qn("w:cs"), FONT)


def style_paragraph(paragraph, before=0, after=0, line=240, exact=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if exact is None:
        fmt.line_spacing = line / 240
        fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    else:
        fmt.line_spacing = Pt(exact)
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    # Keep Word from snapping the line to the document grid or carrying
    # a leftover line onto a new page.
    p_pr = paragraph._p.get_or_add_pPr()
    snap = p_pr.find(qn("w:snapToGrid"))
    if snap is None:
        snap = OxmlElement("w:snapToGrid")
        p_pr.append(snap)
    snap.set(qn("w:val"), "0")
    widow = p_pr.find(qn("w:widowControl"))
    if widow is None:
        widow = OxmlElement("w:widowControl")
        p_pr.append(widow)
    widow.set(qn("w:val"), "0")


def add_text(paragraph, text, size, bold=False):
    run = paragraph.add_run(text)
    set_run_font(run, size, bold)
    return run


def add_hyperlink(paragraph, text, url, size):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_fonts.set(qn("w:cs"), FONT)
    size_el = OxmlElement("w:sz")
    size_el.set(qn("w:val"), str(int(size * 2)))
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), str(int(size * 2)))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4B6E")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([r_fonts, size_el, size_cs, color, underline])
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.set(qn("xml:space"), "preserve")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_heading(doc, text):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, before=4, after=1)
    add_text(paragraph, text.upper(), 11, bold=True)
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1C1F24")
    borders.append(bottom)
    p_pr.append(borders)
    return paragraph


def add_body(doc, text, before=0, after=1):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, before=before, after=after)
    add_text(paragraph, text, 11)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    # Exact leading stops Word from stretching the Symbol bullet font.
    style_paragraph(paragraph, before=0, after=0, exact=13)
    add_text(paragraph, text, 10.5)
    return paragraph


def text_inches(text, size):
    return len(text) * (size * 0.5) / 72


def add_dated_line(doc, left, dates, size, bold, before=0, after=0):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, before=before, after=after)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(7.2), WD_TAB_ALIGNMENT.RIGHT)
    add_text(paragraph, left, size, bold=bold)
    if dates:
        separator = "\t" if text_inches(left, size) + 0.25 + text_inches(dates, size) <= 7.2 else "  "
        add_text(paragraph, f"{separator}{dates}", size)
    return paragraph


def add_role(doc, role):
    title = strip_tags(role.get("title"))
    employer = role.get("employer", "")
    kind = role.get("employer_description", "")
    dates = f"{role.get('start', '')} – {role.get('end', '')}"
    left = f"{title}, {employer}"
    where = f"{employer} · {kind}" if kind else employer
    if text_inches(left, 11) + 0.25 + text_inches(dates, 11) <= 7.2:
        add_dated_line(doc, left, dates, 11, True, before=1)
        if kind:
            add_dated_line(doc, kind, "", 10.5, False)
    elif text_inches(title, 11) + 0.25 + text_inches(dates, 11) <= 7.2:
        add_dated_line(doc, title, dates, 11, True, before=1)
        add_dated_line(doc, where, "", 10.5, False)
    else:
        add_dated_line(doc, title, "", 11, True, before=1)
        add_dated_line(doc, where, dates, 10.5, False)
    bullets = role.get("download_description") or role.get("description") or []
    for item in bullets:
        add_bullet(doc, strip_tags(item))


def prepare_word_layout(doc):
    """Keep Word on the same two pages as the PDF.

    The template snaps every line to an 18pt grid and gives unused paragraphs
    10pt of space after plus 1.15 line spacing. LibreOffice ignores both, so
    the PDF stays at two pages while Word runs to three.
    """
    styles_el = doc.styles.element
    defaults = styles_el.find(qn("w:docDefaults"))
    if defaults is not None:
        for spacing in defaults.findall(".//" + qn("w:spacing")):
            spacing.set(qn("w:before"), "0")
            spacing.set(qn("w:after"), "0")
            spacing.set(qn("w:line"), "240")
            spacing.set(qn("w:lineRule"), "auto")

    normal = doc.styles["Normal"]._element
    p_pr = normal.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = OxmlElement("w:pPr")
        normal.append(p_pr)
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")

    numbering = doc.part.numbering_part._element
    for level in numbering.findall(".//" + qn("w:lvl")):
        style = level.find(qn("w:pStyle"))
        if style is None or style.get(qn("w:val")) != "ListBullet":
            continue
        level_text = level.find(qn("w:lvlText"))
        if level_text is not None:
            level_text.set(qn("w:val"), "•")
        r_pr = level.find(qn("w:rPr"))
        if r_pr is None:
            r_pr = OxmlElement("w:rPr")
            level.append(r_pr)
        r_fonts = r_pr.find(qn("w:rFonts"))
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.insert(0, r_fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            r_fonts.set(qn(attr), FONT)
        for tag, value in (("w:sz", "21"), ("w:szCs", "21")):
            size = r_pr.find(qn(tag))
            if size is None:
                size = OxmlElement(tag)
                r_pr.append(size)
            size.set(qn("w:val"), value)

    grid = doc.sections[0]._sectPr.find(qn("w:docGrid"))
    if grid is not None:
        grid.getparent().remove(grid)


def add_contact(paragraph, email, phone):
    if not email and not phone:
        return
    if email:
        add_text(paragraph, " · ", 11)
        add_hyperlink(paragraph, email, f"mailto:{email}", 11)
    if phone:
        add_text(paragraph, " · ", 11)
        add_text(paragraph, phone, 11)


def contact_from(path):
    contact = {}
    if path:
        loaded = load_yaml(path) or {}
        if not isinstance(loaded, dict):
            raise SystemExit(f"{path} must contain email and phone fields.")
        contact.update(loaded)
    if os.environ.get("RESUME_EMAIL"):
        contact["email"] = os.environ["RESUME_EMAIL"]
    if os.environ.get("RESUME_PHONE"):
        contact["phone"] = os.environ["RESUME_PHONE"]
    email = str(contact.get("email") or "").strip()
    phone = str(contact.get("phone") or "").strip()
    return email, phone


def build(resume, config, email="", phone=""):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    prepare_word_layout(doc)

    name = doc.add_paragraph()
    style_paragraph(name, after=0)
    add_text(name, config.get("name") or config.get("title") or "", 20, bold=True)

    download = resume.get("download") or {}
    headline = download.get("headline")
    if not headline:
        current = (resume.get("roles") or [{}])[0]
        headline = strip_tags(current.get("title"))
    role_line = doc.add_paragraph()
    style_paragraph(role_line, after=0)
    add_text(role_line, headline, 12)

    location = doc.add_paragraph()
    style_paragraph(location, after=1)
    add_text(location, config.get("location", ""), 11)
    add_contact(location, email, phone)

    links = doc.add_paragraph()
    style_paragraph(links, after=2)
    items = list(download.get("links") or config.get("primarylinks") or [])
    if not download.get("links") and config.get("twitter_username"):
        items.append(
            {
                "title": f"@{config['twitter_username']}",
                "url": f"https://twitter.com/{config['twitter_username']}",
            }
        )
    for index, link in enumerate(items):
        if index:
            add_text(links, "   ", 10)
        add_hyperlink(links, link["title"], link["url"], 10)

    add_heading(doc, "Summary")
    for line in str(resume.get("intro_blurb") or "").splitlines():
        line = line.strip()
        if line:
            add_body(doc, line)
    for line in download.get("highlights") or []:
        line = str(line).strip()
        if line:
            add_body(doc, line)

    courses = (resume.get("training") or {}).get("courses") or []
    selected_courses = [course for course in courses if course.get("download")]
    if selected_courses:
        add_heading(doc, "Certifications")
        for course in selected_courses:
            paragraph = doc.add_paragraph()
            style_paragraph(paragraph, before=1, after=1)
            title = course.get("title", "")
            if course.get("link"):
                add_hyperlink(paragraph, title, course["link"], 11)
            else:
                add_text(paragraph, title, 11, bold=True)
            if course.get("date"):
                add_text(paragraph, f"  {course['date']}", 10.5)
    elif courses:
        add_heading(doc, "Training")
        for course in courses:
            paragraph = doc.add_paragraph()
            style_paragraph(paragraph, after=0)
            title = course.get("title", "")
            if course.get("link"):
                add_hyperlink(paragraph, title, course["link"], 11)
            else:
                add_text(paragraph, title, 11, bold=True)
            date = doc.add_paragraph()
            style_paragraph(date, after=3)
            add_text(date, str(course.get("date", "")), 10.5)

    skills = resume.get("skills") or []
    if skills:
        add_heading(doc, "Technologies")
        add_body(doc, " · ".join(skills))

    roles = resume.get("roles") or []
    if roles:
        add_heading(doc, "Work Experience")
        for role in roles:
            add_role(doc, role)

    education = resume.get("education") or []
    if education:
        add_heading(doc, "Education")
        for item in education:
            paragraph = doc.add_paragraph()
            style_paragraph(paragraph, after=0)
            add_text(paragraph, f"{item.get('level', '')}, {item.get('subject', '')}", 11, bold=True)
            detail = doc.add_paragraph()
            style_paragraph(detail, after=2)
            add_text(detail, f"{item.get('school', '')} · {item.get('date', '')}", 11)

    extra = resume.get("additionalinfo")
    if extra and not download.get("highlights"):
        add_heading(doc, "Additional Information")
        for label, body in html_blocks(extra):
            paragraph = doc.add_paragraph()
            style_paragraph(paragraph, after=3)
            if label:
                add_text(paragraph, f"{label}: ", 11, bold=True)
            add_text(paragraph, body, 11)

    core = doc.core_properties
    core.title = f"{config.get('name', '')} resume"
    core.author = config.get("name", "")
    return doc


PUBLIC_NAME = "benjamin-roedell-resume.docx"
PRIVATE_NAME = "benjamin-roedell-resume-private.docx"


def main():
    parser = argparse.ArgumentParser(description="Build a resume docx from the site data.")
    parser.add_argument(
        "--contact",
        type=Path,
        help="Untracked YAML file with email and phone. Omit this for the public resume.",
    )
    parser.add_argument("--docx", type=Path, help="Where to write the Word file.")
    args = parser.parse_args()

    email, phone = contact_from(args.contact)
    private = bool(email or phone)
    if args.docx:
        output = args.docx if args.docx.is_absolute() else Path.cwd() / args.docx
    elif private:
        output = Path.cwd() / PRIVATE_NAME
    else:
        output = ROOT / PUBLIC_NAME
    if private and output.name == PUBLIC_NAME:
        raise SystemExit("Refusing to write contact details into the public resume file.")

    resume = load_yaml(ROOT / "_data" / "resume.yml")
    config = load_yaml(ROOT / "_config.yml")
    document = build(resume, config, email, phone)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(output)


if __name__ == "__main__":
    main()
